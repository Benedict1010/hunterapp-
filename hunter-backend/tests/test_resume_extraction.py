import os
import pytest
import docx
from unittest.mock import MagicMock, patch

from app.core.config import settings
from app.services.resume_extractor import extract_text_from_file, ResumeExtractionError


@pytest.fixture(autouse=True)
def setup_test_upload_dir(monkeypatch, tmp_path):
    test_upload_dir = tmp_path / "test_uploads_extract"
    monkeypatch.setattr(settings, "resume_upload_dir", str(test_upload_dir))
    os.makedirs(test_upload_dir, exist_ok=True)
    return test_upload_dir


def create_auth_header(client, email="extractor_user@example.com"):
    reg_payload = {
        "email": email,
        "password": "password123",
        "full_name": "Extractor User"
    }
    client.post("/auth/register", json=reg_payload)
    login_res = client.post("/auth/login", json={"email": email, "password": "password123"})
    token = login_res.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def test_service_docx_text_extraction(tmp_path):
    # Test Task 4: DOCX text extraction (paragraphs and tables)
    file_path = tmp_path / "test_resume.docx"
    doc = docx.Document()
    doc.add_paragraph("John Doe Resume")
    doc.add_paragraph("Experience: Python Developer")

    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "Advanced"

    doc.save(str(file_path))

    extracted = extract_text_from_file(str(file_path), ".docx")
    assert "John Doe Resume" in extracted
    assert "Experience: Python Developer" in extracted
    assert "Python | Advanced" in extracted


@patch("pypdf.PdfReader")
def test_service_pdf_text_extraction(mock_pdf_reader, tmp_path):
    # Test Task 3: PDF text extraction using mocked PdfReader
    file_path = tmp_path / "test_resume.pdf"
    with open(file_path, "wb") as f:
        f.write(b"%PDF-1.4 dummy contents")

    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Jane Doe Resume\nSkills: FastAPI, PostgreSQL"

    mock_instance = MagicMock()
    mock_instance.pages = [mock_page]
    mock_pdf_reader.return_value = mock_instance

    extracted = extract_text_from_file(str(file_path), ".pdf")
    assert "Jane Doe Resume" in extracted
    assert "Skills: FastAPI, PostgreSQL" in extracted


@patch("pypdf.PdfReader")
def test_empty_or_scanned_pdf_raises_error(mock_pdf_reader, tmp_path):
    # Test Task 3: Scanned/empty PDF handling
    file_path = tmp_path / "scanned.pdf"
    with open(file_path, "wb") as f:
        f.write(b"%PDF-1.4 scanned contents")

    mock_page = MagicMock()
    mock_page.extract_text.return_value = "   "  # Empty whitespace

    mock_instance = MagicMock()
    mock_instance.pages = [mock_page]
    mock_pdf_reader.return_value = mock_instance

    with pytest.raises(ResumeExtractionError) as exc_info:
        extract_text_from_file(str(file_path), ".pdf")
    assert "no extractable text" in str(exc_info.value)


@patch("pypdf.PdfReader")
def test_api_upload_persists_content_text(mock_pdf_reader, client):
    # Test Task 6: Upload triggers extraction and stores text in content_text
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Extracted PDF resume contents."
    mock_instance = MagicMock()
    mock_instance.pages = [mock_page]
    mock_pdf_reader.return_value = mock_instance

    headers = create_auth_header(client, "uploader@example.com")

    response = client.post(
        "/resumes",
        headers=headers,
        files={"file": ("resume.pdf", b"%PDF-1.4 valid", "application/pdf")}
    )
    assert response.status_code == 201
    data = response.json()
    assert data["content_text"] == "Extracted PDF resume contents."


@patch("pypdf.PdfReader")
def test_api_replace_updates_content_text(mock_pdf_reader, client):
    # Test Task 7: Replacement updates content_text
    headers = create_auth_header(client, "replacer@example.com")

    # 1. Initial upload
    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "First content"
    mock_instance1 = MagicMock()
    mock_instance1.pages = [mock_page1]
    mock_pdf_reader.return_value = mock_instance1

    resume = client.post(
        "/resumes",
        headers=headers,
        files={"file": ("resume.pdf", b"%PDF-1.4 first", "application/pdf")}
    ).json()
    assert resume["content_text"] == "First content"

    # 2. Update/Replace
    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = "Updated content"
    mock_instance2 = MagicMock()
    mock_instance2.pages = [mock_page2]
    mock_pdf_reader.return_value = mock_instance2

    updated_res = client.put(
        f"/resumes/{resume['id']}",
        headers=headers,
        files={"file": ("new_resume.pdf", b"%PDF-1.4 updated", "application/pdf")}
    )
    assert updated_res.status_code == 200
    assert updated_res.json()["content_text"] == "Updated content"


@patch("pypdf.PdfReader")
def test_reextract_endpoint_flows_and_access_controls(mock_pdf_reader, client):
    # Test Task 8 & 13: re-extraction endpoint and cross-user isolation
    headers1 = create_auth_header(client, "user_alpha@example.com")
    headers2 = create_auth_header(client, "user_beta@example.com")

    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Alpha Initial Text"
    mock_instance = MagicMock()
    mock_instance.pages = [mock_page]
    mock_pdf_reader.return_value = mock_instance

    resume = client.post(
        "/resumes",
        headers=headers1,
        files={"file": ("alpha.pdf", b"%PDF-1.4 alpha", "application/pdf")}
    ).json()

    # 6. Unauthenticated extraction request is rejected
    res_unauth = client.post(f"/resumes/{resume['id']}/extract")
    assert res_unauth.status_code == 401

    # 7. Another user cannot extract another user's resume
    res_forbidden = client.post(f"/resumes/{resume['id']}/extract", headers=headers2)
    assert res_forbidden.status_code == 403

    # Change mock text to simulate reprocessing updates
    mock_page.extract_text.return_value = "Reprocessed New Text Content"

    # 5. Re-extraction endpoint works for owner
    res_success = client.post(f"/resumes/{resume['id']}/extract", headers=headers1)
    assert res_success.status_code == 200
    assert res_success.json()["content_text"] == "Reprocessed New Text Content"
