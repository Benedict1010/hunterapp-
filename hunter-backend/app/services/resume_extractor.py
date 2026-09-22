import os
import re
import pypdf
import docx

class ResumeExtractionError(Exception):
    pass

def normalize_extracted_text(text: str) -> str:
    # Replace 3 or more consecutive newlines/whitespace lines with just two newlines
    text = re.sub(r'\n\s*\n+', '\n\n', text)
    return text.strip()

def extract_text_from_file(file_path: str, extension: str) -> str:
    if not os.path.exists(file_path):
        raise ResumeExtractionError("Physical resume file not found on disk.")

    ext = extension.lower()
    text = ""

    try:
        if ext == ".pdf":
            reader = pypdf.PdfReader(file_path)
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            text = "\n".join(text_parts)

            if not text.strip():
                raise ResumeExtractionError("PDF file contains no extractable text or consists solely of scanned images.")

        elif ext == ".docx":
            doc = docx.Document(file_path)
            text_parts = []

            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)

            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        text_parts.append(" | ".join(row_cells))

            text = "\n".join(text_parts)

            if not text.strip():
                raise ResumeExtractionError("DOCX file contains no extractable text.")
        else:
            raise ResumeExtractionError(f"Unsupported file extension for extraction: {extension}")

    except ResumeExtractionError:
        raise
    except Exception as e:
        raise ResumeExtractionError(f"Error occurred during text extraction: {str(e)}")

    return normalize_extracted_text(text)
